from datetime import datetime, date
from openpyxl import load_workbook
from app.logger import log


def _fmt(value) -> str:
    """Convert a cell value to clean text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        # Drop the time part when it's midnight (pure dates)
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 5.0 -> 5
    return str(value).strip()


def parse_xlsx(file_path: str) -> str:
    """
    Parses Excel files (.xlsx, .xlsm) into row-wise text for RAG.
    Each sheet gets a header, and each row becomes:
    'col1: value | col2: value | ...' so every chunk keeps its column context.
    The first non-empty row of each sheet is treated as the header row.
    """
    with log.span(f"📄 Excel Parsing {file_path}"):
        wb = None
        try:
            # read_only = low memory for big files
            # data_only = read cached formula results instead of formula strings
            wb = load_workbook(file_path, read_only=True, data_only=True)

            sections: list[str] = []
            total_rows = 0

            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    log.info(f"Skipping hidden sheet '{ws.title}'")
                    continue

                headers: list[str] | None = None
                rows: list[str] = []

                for raw_row in ws.iter_rows(values_only=True):
                    cells = [_fmt(c) for c in raw_row]

                    if not any(cells):  # skip fully empty rows
                        continue

                    # First non-empty row = headers
                    if headers is None:
                        headers = [
                            h if h else f"Column {i + 1}" for i, h in enumerate(cells)
                        ]
                        continue

                    parts = [
                        f"{headers[i] if i < len(headers) else f'Column {i + 1}'}: {val}"
                        for i, val in enumerate(cells)
                        if val
                    ]
                    if parts:
                        rows.append(" | ".join(parts))

                if rows:
                    sections.append(f"## Sheet: {ws.title}\n" + "\n".join(rows))
                    total_rows += len(rows)
                else:
                    log.info(f"Sheet '{ws.title}' has no data rows")

            text = "\n\n".join(sections)

            if not text.strip():
                log.warning(
                    f"⚠️ No content extracted from {file_path}. "
                    "If it contains formulas, they may have no cached values."
                )
            else:
                log.info(f"✅ Parsed {total_rows} rows across {len(sections)} sheets, {len(text)} characters")

            return text
        except Exception as e:
            log.error(f"❌ Excel Parse Failed: {e}")
            raise e
        finally:
            if wb is not None:
                wb.close()